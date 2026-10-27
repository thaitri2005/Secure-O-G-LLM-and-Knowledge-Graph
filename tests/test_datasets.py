import json
from pathlib import Path
import tempfile
import unicodedata
import unittest
import zipfile

from graph_builder.datasets.importer import run_import, validate_bundle
from graph_builder.datasets.readers import iter_rows
from graph_builder.publication import current_release
from graph_builder.datasets.infer import infer_mapping
from graph_builder.datasets.mapping import validate_mapping
from graph_builder.datasets.values import ValueProblem, parse_value

WELLS = "ma_gieng;ten;do_sau_m;ngay_khoan;mo;vi_tri\n" \
        "G001;Giếng Một;1955;24/10/2016;Bạch Hổ;AB0827\n" \
        "G002;Giếng Hai;#;01/02/2017;Bạch Hổ;AB0012\n" \
        "G003;Giếng Ba;2100;15/03/2018;;CD0003\n"
TESTS = "test_id,ma_gieng,ket_qua\nT1,G001,ok\nT2,G999,fail\n"
TONES = "̣̀́̃̉"


def cp1258(text: str) -> bytes:
    """Encode as Windows-1258 does: base letters precomposed, tone marks as combining characters."""
    out = []
    for char in text:
        try:
            out.append(char.encode("cp1258"))
        except UnicodeEncodeError:
            decomposed = unicodedata.normalize("NFD", char)
            base = unicodedata.normalize("NFC", "".join(c for c in decomposed if c not in TONES))
            out.append((base + "".join(c for c in decomposed if c in TONES)).encode("cp1258"))
    return b"".join(out)


def mapping(base: Path, max_errors: int = 0) -> dict:
    value = {
        "dataset_id": "demo", "title": "Demo", "license": {"id": "test", "url": ""},
        "sources": [
            {"source_id": "wells", "table": "well", "path": "wells.csv", "constants": {"region": "vn"}},
            {"source_id": "tests", "table": "test", "path": "tests.zip", "member": "tests.csv"}],
        "tables": {
            "well": {"layout": {"type": "csv", "delimiter": ";", "encoding": "auto"}, "withheld_chars": ["#"],
                     "id": "demo:well:{well_code}", "key": ["well_code"],
                     "columns": [{"name": "well_code", "source": "ma_gieng", "type": "string", "required": True},
                                 {"name": "name", "source": "ten", "type": "string"},
                                 {"name": "depth_m", "source": "do_sau_m", "type": "integer"},
                                 {"name": "spud_date", "source": "ngay_khoan", "type": "date", "formats": ["%d/%m/%Y"]},
                                 {"name": "field", "source": "mo", "type": "string"},
                                 {"source": "vi_tri", "split": [{"name": "area", "width": 2, "type": "string"},
                                                                 {"name": "block", "width": 4, "type": "string"}]}],
                     "derived": {"block_key": {"op": "strip_leading_zeros", "from": "block"}},
                     "checks": [{"check": "range", "name": "depth_range", "severity": "error",
                                 "fields": ["depth_m"], "min": 0, "max": 12000}]},
            "test": {"layout": {"type": "csv"}, "id": "demo:test:{test_id}", "key": ["test_id"],
                     "columns": [{"name": "test_id", "type": "string"}, {"name": "ma_gieng", "type": "string"},
                                 {"name": "ket_qua", "type": "string", "values": ["ok", "fail"]}],
                     "checks": [{"check": "reference", "name": "unknown_well", "severity": "warning",
                                 "field": "ma_gieng", "table": "well"}]}},
        "graph": {
            "nodes": [{"label": "Well", "table": "well", "id": "well:{well_code}", "properties": ["name", "depth_m"]},
                      {"label": "Field", "table": "well", "id": "field:{field}", "properties": {"name": "field"}},
                      {"label": "Block", "table": "well", "id": "block:{area}-{block_key}"}],
            "edges": [{"type": "IN_FIELD", "table": "well", "from": "well:{well_code}", "to": "field:{field}"},
                      {"type": "IN_BLOCK", "table": "well", "from": "well:{well_code}", "to": "block:{area}-{block_key}"},
                      {"type": "TESTED", "table": "test", "from": "well:{ma_gieng}", "to": "well:{ma_gieng}",
                       "require_existing": True}]},
        "error_policy": {"max_errors": max_errors}, "_base": str(base)}
    validate_mapping(value)
    return value


class DatasetImportTest(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.base = Path(self.directory.name)
        (self.base / "wells.csv").write_bytes(cp1258(WELLS))
        with zipfile.ZipFile(self.base / "tests.zip", "w") as archive:
            archive.writestr("tests.csv", TESTS)

    def tearDown(self):
        self.directory.cleanup()

    def test_import_types_withheld_graph_and_provenance(self):
        output = self.base / "out"
        summary = run_import(mapping(self.base), self.base, output)
        self.assertEqual(summary["records"], {"test": 2, "well": 3})
        self.assertEqual(summary["nodes"], {"Block": 3, "Field": 1, "Well": 3})
        self.assertEqual(summary["edges"], {"IN_BLOCK": 3, "IN_FIELD": 2, "TESTED": 1})
        self.assertEqual(summary["warnings"]["unknown_well"]["count"], 1)
        release = current_release(output)
        self.assertEqual(str(release), summary["release"])
        wells = [json.loads(line) for line in (release / "records/well.jsonl").read_text(encoding="utf-8").splitlines()]
        self.assertEqual(wells[0]["name"], "Giếng Một")  # cp1258 decoded
        self.assertEqual((wells[0]["spud_date"], wells[0]["block_key"], wells[0]["region"]), ("2016-10-24", "827", "vn"))
        self.assertEqual((wells[1]["depth_m"], wells[1]["withheld"]), (None, ["depth_m"]))
        self.assertEqual(wells[2]["source"]["row"], 3)
        self.assertEqual(validate_bundle(output), [])

    def test_errors_block_output_unless_budgeted(self):
        (self.base / "wells.csv").write_text(WELLS + "G001;Trùng;10;01/01/2020;X;AB0001\n", encoding="utf-8")
        output = self.base / "out"
        with self.assertRaises(ValueError):
            run_import(mapping(self.base), self.base, output)
        self.assertFalse(output.exists())
        summary = run_import(mapping(self.base, max_errors=1), self.base, output)
        self.assertEqual(summary["errors"]["duplicate_key"]["count"], 1)
        self.assertEqual(len((current_release(output) / "rejects.jsonl").read_text(encoding="utf-8").splitlines()), 1)

    def test_failed_reimport_keeps_the_current_release(self):
        output = self.base / "out"
        first = run_import(mapping(self.base), self.base, output)["release"]
        (self.base / "wells.csv").write_text(WELLS + "G001;Trùng;10;01/01/2020;X;AB0001\n", encoding="utf-8")
        with self.assertRaises(ValueError):
            run_import(mapping(self.base), self.base, output)
        self.assertEqual(str(current_release(output)), first)
        self.assertEqual(validate_bundle(output), [])
        self.assertEqual([p.name for p in self.base.iterdir() if p.name.startswith(".out-")], [])  # staging removed

    def test_duplicate_record_ids_fail_without_a_natural_key(self):
        """E9 probe: two identical generated IDs used to import, then fail validation after replacing the bundle."""
        (self.base / "codes.csv").write_text("code\nA\nA\n", encoding="utf-8")
        value = {"dataset_id": "d", "title": "d", "license": {"id": "t", "url": ""},
                 "sources": [{"source_id": "c", "table": "c", "path": "codes.csv"}],
                 "tables": {"c": {"layout": {"type": "csv"}, "id": "d:{code}",
                                  "columns": [{"name": "code", "type": "string"}]}},
                 "_base": str(self.base)}
        validate_mapping(value)
        with self.assertRaisesRegex(ValueError, "duplicate_record_id"):
            run_import(value, self.base, self.base / "codes")
        self.assertFalse((self.base / "codes").exists())

    def test_duplicate_headers_are_refused(self):
        """E9 probe: code,depth,depth with A,100,999 used to keep only depth=999."""
        (self.base / "dup.csv").write_text("code,depth,depth\nA,100,999\n", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "duplicate column names"):
            list(iter_rows(self.base / "dup.csv", None, {"type": "csv"}))
        (self.base / "blank.csv").write_text("code,,\nA,1,2\n", encoding="utf-8")
        self.assertEqual(list(iter_rows(self.base / "blank.csv", None, {"type": "csv"})),
                         [(1, {"code": "A", "_column2": "1", "_column3": "2"})])

    def test_pinned_digest_mismatch_fails(self):
        value = mapping(self.base)
        value["sources"][0]["sha256"] = "0" * 64
        with self.assertRaises(ValueError):
            run_import(value, self.base, self.base / "out")

    def test_values(self):
        self.assertEqual(parse_value({"type": "decimal", "decimal_mark": ",", "thousands": "."}, "1.234,5"), 1234.5)
        with self.assertRaises(ValueProblem):
            parse_value({"type": "date", "formats": ["%Y%m%d"]}, "2016101")
        with self.assertRaises(ValueProblem):
            parse_value({"type": "integer"}, "12a")

    def test_infer_drafts_a_mapping(self):
        draft = infer_mapping(self.base / "wells.csv")
        table = draft["tables"]["wells"]
        types = {column["name"]: column["type"] for column in table["columns"]}
        self.assertEqual(table["layout"]["delimiter"], ";")
        self.assertEqual(types["ngay_khoan"], "date")
        self.assertEqual(types["vi_tri"], "string")
        self.assertEqual(table["key"], ["ma_gieng"])

    def test_xlsx(self):
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.append(["code", "depth"])
        sheet.append(["A", 10])
        sheet.append(["B", 12.0])
        book.save(self.base / "w.xlsx")
        value = {"dataset_id": "x", "title": "x", "license": {"id": "t", "url": ""},
                 "sources": [{"source_id": "w", "table": "w", "path": "w.xlsx"}],
                 "tables": {"w": {"layout": {"type": "xlsx"}, "id": "x:{code}", "key": ["code"],
                                  "columns": [{"name": "code", "type": "string"}, {"name": "depth", "type": "integer"}]}},
                 "_base": str(self.base)}
        validate_mapping(value)
        summary = run_import(value, self.base, self.base / "xo")
        self.assertEqual(summary["records"], {"w": 2})


if __name__ == "__main__":
    unittest.main()
